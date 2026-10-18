"""Reusable python-pptx helpers for editing a template deck in place.

Import from your build script:

    import sys; sys.path.insert(0, "<skill_dir>/scripts")
    from pptx_helpers import *

Target shapes by shape_id or name (see inspect_template.py output), never by a
cached index: deleting a shape shifts the index of every later shape in
slide.shapes, original template shapes included. If you must work by index,
look up every shape you need first, then delete in descending index order.

Extracted from a working template-edit build (DPPA 2026 deck); fonts and
colours are defaults, so pass the template's own values.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT_HEAD = "Calibri"
FONT_BODY = "Calibri"
TEXT = "131314"
MUTED = "5B646E"
ACCENT = "1B786E"


def rgb(hex_color):
    """'1B786E' or '#1B786E' -> RGBColor."""
    hex_color = hex_color.strip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


# --- finding shapes -------------------------------------------------------

def shape_by_id(slide, shape_id):
    """Return the shape with this shape_id (stable across deletions)."""
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    raise KeyError(f"no shape with id {shape_id} on this slide")


def shape_by_name(slide, name):
    """Return the single shape with this name; raise if missing or ambiguous."""
    hits = [sh for sh in slide.shapes if sh.name == name]
    if len(hits) != 1:
        raise KeyError(f"{len(hits)} shapes named {name!r} on this slide")
    return hits[0]


def delete_shape(shape):
    """Remove a shape from its slide. Later indices in slide.shapes shift down by one."""
    el = shape._element
    el.getparent().remove(el)


def delete_shapes(shapes):
    """Delete several shapes, resolved up front, so index shifts can't hit the wrong one."""
    for sh in list(shapes):
        delete_shape(sh)


# --- text -----------------------------------------------------------------

def suppress_bullet(paragraph):
    """Drop any inherited bullet so a reused list box doesn't render a stray bullet."""
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def clear_text(shape):
    tf = shape.text_frame
    tf.clear()
    return tf


def set_text(shape, text, size=14, bold=False, color=TEXT, font=FONT_BODY, align=None):
    """Replace a shape's text with one run: zero margins, no bullet, given font."""
    tf = clear_text(shape)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    suppress_bullet(p)
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(color)


def add_textbox(slide, x, y, w, h, text, size=10, bold=False, color=TEXT, font=FONT_BODY, align=None):
    """New text box at (x, y, w, h) inches, styled via set_text."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(box, text, size=size, bold=bold, color=color, font=font, align=align)
    return box


def add_bullets(slide, x, y, w, h, items, size=9.5, color=TEXT, font=FONT_BODY, space_after=5):
    """New text box with one paragraph per item."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = font
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(space_after)
    return box


# --- panels and pictures --------------------------------------------------

def cover(slide, x, y, w, h, color="FFFFFF"):
    """Solid rectangle, e.g. to hide an old dark panel before placing new content."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.color.rgb = rgb(color)
    shape.shadow.inherit = False
    return shape


def add_stat_card(slide, x, y, w, h, value, label, accent=ACCENT,
                  font_head=FONT_HEAD, font_body=FONT_BODY, muted=MUTED):
    """Light card with an accent top rule, a big value and a small label."""
    from pptx.enum.text import PP_ALIGN
    card = cover(slide, x, y, w, h, "FBFCFC")
    card.line.color.rgb = rgb("E2EAE9")
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.045))
    rule.fill.solid()
    rule.fill.fore_color.rgb = rgb(accent)
    rule.line.fill.background()
    rule.shadow.inherit = False
    add_textbox(slide, x + 0.08, y + 0.15, w - 0.16, 0.27, value, 18, True, accent, font_head, PP_ALIGN.CENTER)
    add_textbox(slide, x + 0.08, y + 0.5, w - 0.16, 0.24, label, 7.2, False, muted, font_body, PP_ALIGN.CENTER)
    return card


def add_picture_cover(slide, path, x, y, w, h):
    """Place a PNG at (x, y, w, h) inches, normally where the deleted picture sat."""
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def replace_picture(slide, old_picture, png_path):
    """Delete a picture and put a new PNG at exactly the same position and size."""
    x, y, w, h = old_picture.left, old_picture.top, old_picture.width, old_picture.height
    delete_shape(old_picture)
    return slide.shapes.add_picture(str(png_path), x, y, width=w, height=h)
